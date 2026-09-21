"""SuperTonic — servidor HTTP (FastAPI) com UI web, fila de jobs e API key opcional.

Rotas próprias (além de /v1/* do ``supertonic serve``):
  GET  /                      interface web (static/)
  GET  /health                healthcheck + capacidades
  GET  /api/voices            vozes, formatos e idioma padrão
  POST /api/jobs              cria um job (texto | arquivo | link) → {"id"}
  GET  /api/jobs/{id}         estado/progresso do job
  GET  /api/jobs/{id}/chunks/{n}  áudio WAV do bloco n (reprodução progressiva)
  GET  /api/jobs/{id}/audio   arquivo final (mp3/wav/ogg/flac)
  POST /usar                  rota síncrona (compatibilidade)
"""

from __future__ import annotations

import asyncio
import hashlib
from contextlib import asynccontextmanager
from importlib.metadata import version, PackageNotFoundError
import html
import io
import json
import logging
import os
import re
import secrets
import shutil
import subprocess
import tempfile
import time
import threading
import urllib.parse
import uuid
from pathlib import Path
from typing import Any, Dict, List, Optional

import numpy as np
import soundfile as sf
import uvicorn
from fastapi import FastAPI, File, Form, Request, UploadFile
from fastapi.responses import FileResponse, JSONResponse, Response
from fastapi.staticfiles import StaticFiles
from starlette.middleware.base import BaseHTTPMiddleware

from supertonic.server.app import create_app
from supertonic.server.audio import SUPPORTED_FORMATS, encode_audio, format_to_mime
from supertonic.server.routes import UnknownVoice, _do_synthesize

from backend_resources import BodyLimitMiddleware, Cancelled, ResourceError, save_upload, trusted_peer, run_process
import supabase_persistence

from textnorm import clean_layout, normalize_for_tts, split_chunks, strip_repeated_lines

logger = logging.getLogger("supertonic.ui")

BASE_DIR = Path(__file__).resolve().parent
STATIC_DIR = BASE_DIR / "static"
JOBS_DIR = Path(os.environ.get("JOBS_DIR", tempfile.gettempdir())) / "supertonic-jobs"
JOBS_DIR.mkdir(parents=True, exist_ok=True)

# ---------------------------------------------------------------------------
# Configuração
# ---------------------------------------------------------------------------
MODEL = os.environ.get("SUPERSONIC_MODEL", os.environ.get("SUPERTONIC_MODEL", "supertonic-3"))
QUEUE_MAX_SIZE = max(1, int(os.environ.get("QUEUE_MAX_SIZE", "16")))
TRUSTED_PROXIES = os.environ.get("TRUSTED_PROXIES", "")
ALLOW_REMOTE_DOWNLOADS = os.environ.get("ALLOW_REMOTE_DOWNLOADS", "0") == "1"
CACHE_REVISION = os.environ.get("CACHE_REVISION", "1")
try:
    ENGINE_VERSION = version("supertonic")
except PackageNotFoundError:
    ENGINE_VERSION = "stub"
DEFAULT_LANG = os.environ.get("DEFAULT_LANG", "pt")
MAX_TEXT_CHARS = int(os.environ.get("MAX_TEXT_CHARS", "120000"))
MAX_UPLOAD_MB = int(os.environ.get("MAX_UPLOAD_MB", "200"))
JOB_TTL_SECONDS = int(os.environ.get("JOB_TTL_SECONDS", str(3 * 3600)))
CHUNK_GAP_SECONDS = float(os.environ.get("CHUNK_GAP_SECONDS", "0.25"))
DEFAULT_PARAGRAPH_PAUSE = float(os.environ.get("PARAGRAPH_PAUSE", "0.6"))
RATE_LIMIT_PER_HOUR = int(os.environ.get("RATE_LIMIT_PER_HOUR", "40"))
CACHE_TTL_SECONDS = int(os.environ.get("CACHE_TTL_SECONDS", str(7 * 24 * 3600)))
SHARE_TTL_SECONDS = int(os.environ.get("SHARE_TTL_SECONDS", str(24 * 3600)))
CACHE_DIR = JOBS_DIR / "_cache"
SHARE_DIR = JOBS_DIR / "_share"
CACHE_DIR.mkdir(parents=True, exist_ok=True)
SHARE_DIR.mkdir(parents=True, exist_ok=True)
MP3_BITRATE = os.environ.get("MP3_BITRATE", "64k")

WHISPER_MODEL = os.environ.get("WHISPER_MODEL", "small")
WHISPER_DEVICE = os.environ.get("WHISPER_DEVICE", "cpu")
OPENAI_KEY = (os.environ.get("OPENAI_API_KEY") or os.environ.get("openai_key") or os.environ.get("OPENAI_KEY") or "").strip()
OPENAI_TRANSCRIBE_MODEL = os.environ.get("OPENAI_TRANSCRIBE_MODEL", "gpt-4o-mini-transcribe")
OPENAI_MAX_UPLOAD = 24 * 1024 * 1024  # limite da API é 25 MB

PUBLIC_EXACT = {"/", "/usar", "/api/voices", "/manifest.webmanifest", "/sw.js", "/favicon.svg", "/favicon.ico", "/api/feed.xml", "/feed.xml"}
PUBLIC_PREFIXES = ("/health", "/docs", "/redoc", "/openapi.json", "/static", "/api/jobs", "/s", "/api/documents")

TEXT_EXT = {".txt", ".md", ".markdown", ".csv", ".json", ".html", ".htm", ".srt", ".vtt"}
PDF_EXT = {".pdf"}
DOCX_EXT = {".docx"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tiff"}
MEDIA_EXT = {".mp4", ".webm", ".mov", ".mkv", ".avi", ".mp3", ".wav", ".m4a", ".ogg", ".flac", ".aac", ".opus", ".oga"}


def _has(module: str) -> bool:
    import importlib.util

    return importlib.util.find_spec(module) is not None


def _ffmpeg() -> Optional[str]:
    exe = shutil.which("ffmpeg")
    if exe:
        return exe
    try:  # binário estático para desenvolvimento local
        import imageio_ffmpeg

        return imageio_ffmpeg.get_ffmpeg_exe()
    except Exception:
        return None


FFMPEG = _ffmpeg()
OUTPUT_FORMATS = list(SUPPORTED_FORMATS) + (["mp3", "m4b"] if FFMPEG else [])


# ---------------------------------------------------------------------------
# Erros
# ---------------------------------------------------------------------------
class UsarError(Exception):
    def __init__(self, message: str, status: int = 400, code: str = "invalid_request"):
        super().__init__(message)
        self.message, self.status, self.code = message, status, code


def _err(status: int, message: str, code: str, type_: str = "invalid_request_error") -> JSONResponse:
    return JSONResponse(status_code=status, content={"error": {"message": message, "type": type_, "code": code}})


class ApiKeyMiddleware(BaseHTTPMiddleware):
    def __init__(self, app, api_key: str):
        super().__init__(app)
        self.api_key = api_key

    async def dispatch(self, request: Request, call_next) -> Response:
        path = request.url.path.rstrip("/") or "/"
        if path in PUBLIC_EXACT or any(path == p or path.startswith(p + "/") for p in PUBLIC_PREFIXES):
            return await call_next(request)
        provided = request.headers.get("x-api-key") or ""
        auth = request.headers.get("authorization") or ""
        if auth.lower().startswith("bearer "):
            provided = auth[7:].strip()
        if not provided or not secrets.compare_digest(provided, self.api_key):
            return _err(401, "Chave inválida. Envie o cabeçalho X-API-Key ou Authorization: Bearer <chave>.",
                        "invalid_api_key", "authentication_error")
        return await call_next(request)


# ---------------------------------------------------------------------------
# Extração de texto
# ---------------------------------------------------------------------------
def parse_pages(spec: Optional[str], total: int) -> Optional[List[int]]:
    """'3-10, 12' → [2..9, 11] (índices zero-based, limitados a ``total``)."""
    if not spec or not spec.strip():
        return None
    if len(spec) > 2048:
        raise UsarError("Seleção de páginas muito longa.", 400, "bad_pages")
    if total > 10000:
        raise UsarError("PDF excede 10000 páginas.", 400, "bad_pages")
    out: set = set()
    for part in re.split(r"[,;\s]+", spec.strip()):
        if not part:
            continue
        m = re.fullmatch(r"(\d{1,12})(?:-(\d{1,12}))?", part)
        if not m:
            raise UsarError(f"Intervalo de páginas inválido: '{part}'. Use algo como 3-10, 12.", 400, "bad_pages")
        a = int(m.group(1)); b = int(m.group(2) or a)
        if a > b:
            a, b = b, a
        for n in range(max(1, a), min(total, b) + 1):
            if 1 <= n <= total:
                out.add(n - 1)
    if not out:
        raise UsarError(f"Nenhuma página válida (o PDF tem {total}).", 400, "bad_pages")
    return sorted(out)


def _extract_pdf(data: Path, pages_spec: Optional[str] = None) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:  # pragma: no cover
        raise UsarError("Leitura de PDF indisponível (pypdf).", 501, "pdf_unavailable") from e
    reader = PdfReader(str(data))
    if len(reader.pages) > 10000:
        raise UsarError("PDF excede 10000 páginas.", 400, "bad_pages")
    wanted = parse_pages(pages_spec, len(reader.pages))
    selected = [reader.pages[i] for i in wanted] if wanted is not None else list(reader.pages)
    pages: List[str] = []
    for page in selected:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    pages = strip_repeated_lines(pages)
    text = clean_layout("\n\n".join(p for p in pages if p.strip()))
    if not text:
        raise UsarError("Esse PDF parece ser só imagem (escaneado). Não consegui extrair texto.", 422, "pdf_no_text")
    return text


def _extract_docx(data: Path) -> str:
    try:
        import docx
    except ImportError as e:  # pragma: no cover
        raise UsarError("Leitura de Word indisponível (python-docx).", 501, "docx_unavailable") from e
    d = docx.Document(str(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return clean_layout("\n\n".join(parts))


def _extract_plain(data: Path, ext: str) -> str:
    with open(data, "rb") as source:
        text = source.read(MAX_TEXT_CHARS * 4 + 1).decode("utf-8", errors="replace")
    if ext in {".html", ".htm"}:
        text = re.sub(r"<script.*?</script>|<style.*?</style>", " ", text, flags=re.S | re.I)
        text = re.sub(r"<(br|/p|/div|/h\d|/li)[^>]*>", "\n\n", text, flags=re.I)
        text = re.sub(r"<[^>]+>", " ", text)
    if ext in {".srt", ".vtt"}:
        text = re.sub(r"^\d+\s*$", "", text, flags=re.M)
        text = re.sub(r"\d{2}:\d{2}:\d{2}[.,]\d{3} --> .*$", "", text, flags=re.M)
        text = text.replace("WEBVTT", "")
    return clean_layout(text)


def _extract_image(data: Path) -> str:
    try:
        import pytesseract
        from PIL import Image
    except ImportError as e:
        raise UsarError("OCR de imagem não está instalado neste servidor.", 501, "ocr_unavailable") from e
    if not shutil.which("tesseract"):
        raise UsarError("Tesseract não encontrado no servidor.", 501, "ocr_unavailable")
    img = Image.open(data)
    text = pytesseract.image_to_string(img, lang=os.environ.get("TESSERACT_LANG", "por+eng"))
    text = clean_layout(text)
    if not text:
        raise UsarError("Não encontrei texto legível na imagem.", 422, "ocr_no_text")
    return text


# ---------------------------------------------------------------------------
# Transcrição: OpenAI (rápida) → faster-whisper local (fallback)
# ---------------------------------------------------------------------------
_whisper_model = None


def _load_whisper():
    global _whisper_model
    if _whisper_model is None:
        try:
            from faster_whisper import WhisperModel
        except ImportError as e:
            raise UsarError("Transcrição de áudio/vídeo não está habilitada neste servidor.", 501, "transcription_unavailable") from e
        compute = "int8" if WHISPER_DEVICE == "cpu" else "float16"
        logger.info("Carregando faster-whisper %s (%s)…", WHISPER_MODEL, WHISPER_DEVICE)
        _whisper_model = WhisperModel(WHISPER_MODEL, device=WHISPER_DEVICE, compute_type=compute)
    return _whisper_model


def _to_compact_audio(path: str, workdir: str, check_cancel=lambda: None) -> str:
    """Extrai/compacta o áudio (mono, 16 kHz, mp3 32k) para caber na API e ser rápido."""
    if not FFMPEG:
        return path
    out = os.path.join(workdir, "audio16k.mp3")
    cmd = [FFMPEG, "-y", "-loglevel", "error", "-protocol_whitelist", "file,pipe", "-i", path, "-vn", "-ac", "1", "-ar", "16000", "-b:a", "32k", out]
    run_process(cmd, check_cancel, timeout=1800)
    return out


def _transcribe_openai(path: str, lang: Optional[str], progress) -> str:
    import requests

    size = os.path.getsize(path)
    if size > OPENAI_MAX_UPLOAD:
        raise UsarError("Áudio grande demais para a API de transcrição.", 413, "audio_too_large")
    progress("transcribe", 20, "Transcrevendo com IA…")
    data = {"model": OPENAI_TRANSCRIBE_MODEL, "response_format": "json"}
    if lang and lang != "na":
        data["language"] = lang
    with open(path, "rb") as f:
        r = requests.post(
            "https://api.openai.com/v1/audio/transcriptions",
            headers={"Authorization": f"Bearer {OPENAI_KEY}"},
            data=data,
            files={"file": (os.path.basename(path), f, "audio/mpeg")},
            timeout=600,
        )
    if r.status_code >= 400:
        raise RuntimeError(f"OpenAI {r.status_code}: {r.text[:200]}")
    return (r.json().get("text") or "").strip()


def _transcribe_local(path: str, lang: Optional[str], progress) -> str:
    progress("transcribe", 15, "Carregando transcritor…")
    model = _load_whisper()
    progress("transcribe", 15, "Transcrevendo no servidor (pode demorar)…")
    segments, info = model.transcribe(path, vad_filter=True, beam_size=1, language=None if lang in (None, "na") else lang)
    total = float(getattr(info, "duration", 0) or 0)
    out: List[str] = []
    for s in segments:
        if s.text.strip():
            out.append(s.text.strip())
        progress("transcribe", 15 + (min(40, int(40 * s.end / total)) if total else 0), "Transcrevendo…")
    return " ".join(out)


def _transcribe(path: str, lang: Optional[str], progress) -> str:
    with tempfile.TemporaryDirectory() as tmp:
        try:
            src = _to_compact_audio(path, tmp, getattr(getattr(progress, "__self__", None), "check_cancel", lambda: None))
        except Cancelled:
            raise
        except Exception as e:
            logger.warning("ffmpeg falhou (%s); usando arquivo original", e)
            src = path
        if OPENAI_KEY:
            try:
                text = _transcribe_openai(src, lang, progress)
                if text:
                    return clean_layout(text)
            except (UsarError, Cancelled):
                raise
            except Exception as e:
                logger.warning("OpenAI transcription falhou, usando local: %s", e)
        return clean_layout(_transcribe_local(src, lang, progress))


def _download_media(url: str, workdir: str, progress) -> str:
    if not ALLOW_REMOTE_DOWNLOADS:
        raise UsarError("Downloads remotos desabilitados. Envie o arquivo.", 403, "remote_download_disabled")
    parsed = urllib.parse.urlparse(url)
    if parsed.scheme not in {"http", "https"}:
        raise UsarError("Link inválido. Use http(s).", 400, "bad_url")
    progress("download", 5, "Baixando o vídeo…")
    ytdlp = shutil.which("yt-dlp")
    if ytdlp:
        cmd = [ytdlp]
    elif _has("yt_dlp"):
        cmd = ["python", "-m", "yt_dlp"]
    else:
        cmd = None
    out_tpl = os.path.join(workdir, "media.%(ext)s")
    if cmd:
        cmd += ["-f", "bestaudio/best", "--no-playlist", "--quiet", "--no-warnings",
                "--max-filesize", f"{MAX_UPLOAD_MB * 4}m", "-o", out_tpl, url]
        proc = subprocess.run(cmd, capture_output=True, text=True, timeout=900)
        if proc.returncode != 0:
            last = ((proc.stderr or "").strip().splitlines() or ["falha desconhecida"])[-1]
            raise UsarError("Não consegui baixar o vídeo (" + last[:160] + "). Baixe o arquivo e envie na aba Arquivo.",
                            502, "download_failed")
    else:
        import urllib.request

        with urllib.request.urlopen(url, timeout=120) as r, open(os.path.join(workdir, "media.bin"), "wb") as f:
            shutil.copyfileobj(r, f, length=1 << 20)
    files = [p for p in Path(workdir).iterdir() if p.name.startswith("media.")]
    if not files:
        raise UsarError("Download vazio.", 502, "download_failed")
    return str(files[0])


def _extract_from_upload(filename: str, data: Path, lang: Optional[str], progress, pages: Optional[str] = None) -> str:
    ext = Path(filename or "").suffix.lower()
    progress("read", 8, "Lendo o documento…")
    if ext in PDF_EXT:
        return _extract_pdf(data, pages)
    if ext in DOCX_EXT:
        return _extract_docx(data)
    if ext == ".epub":
        from document_features import extract_epub, EpubLimits, DocumentFeatureError
        try:
            if data.stat().st_size > EpubLimits().max_archive_bytes:
                raise UsarError("EPUB excede o limite de 64 MiB.", 413, "epub_too_large")
            return extract_epub(data.read_bytes())
        except DocumentFeatureError as exc:
            raise UsarError(str(exc), 422, "invalid_epub") from exc
    if ext in TEXT_EXT or ext == "":
        return _extract_plain(data, ext)
    if ext in IMAGE_EXT:
        progress("read", 10, "Reconhecendo texto na imagem…")
        return _extract_image(data)
    if ext in MEDIA_EXT:
        return _transcribe(str(data), lang, progress)
    raise UsarError(f"Tipo de arquivo não suportado: {ext or 'sem extensão'}.", 415, "unsupported_type")


# ---------------------------------------------------------------------------
# Jobs
# ---------------------------------------------------------------------------
class Job:
    def __init__(self, **kw):
        self.id = uuid.uuid4().hex[:12]
        self.created = time.time()
        self.status = "queued"          # queued | running | done | error
        self.stage = "queued"           # queued | download | read | transcribe | tts | encode | done
        self.percent = 0
        self.message = "Na fila…"
        self.error: Optional[Dict[str, Any]] = None
        self.text = ""
        self.chunks: List[str] = []
        self.chunks_ready = 0
        self.duration = 0.0
        self.final: Optional[str] = None
        self.format = kw.get("format", "wav")
        self.voice = kw.get("voice", "F1")
        self.lang = kw.get("lang") or DEFAULT_LANG
        self.speed = kw.get("speed")
        self.pause = kw.get("pause", DEFAULT_PARAGRAPH_PAUSE)
        self.pages = kw.get("pages")
        mode_val = kw.get("mode", "normal")
        self.mode = str(getattr(mode_val, "default", mode_val) if hasattr(mode_val, "default") else (mode_val or "normal"))
        trans_val = kw.get("translate", False)
        self.translate = bool(getattr(trans_val, "default", trans_val) if hasattr(trans_val, "default") else trans_val)
        skip_val = kw.get("smart_skip", False)
        self.smart_skip = bool(getattr(skip_val, "default", skip_val) if hasattr(skip_val, "default") else skip_val)
        self.podcast_voices: List[str] = []
        self.cached = False
        self.input = kw  # text | file(name, data) | url
        self.dir = JOBS_DIR / self.id
        self.dir.mkdir(parents=True, exist_ok=True)
        self.truncated = False
        self.cancel_event = threading.Event()
        self.finished = None
        self.completion = asyncio.Event()

    def check_cancel(self):
        if self.cancel_event.is_set():
            raise Cancelled()

    def progress(self, stage: str, percent: int, message: str):
        self.check_cancel()
        self.stage, self.percent, self.message = stage, max(self.percent, min(99, percent)), message

    def public(self, request: Request) -> Dict[str, Any]:
        base = f"/api/jobs/{self.id}"
        return {
            "id": self.id, "status": self.status, "stage": self.stage, "percent": self.percent,
            "message": self.message, "error": self.error, "text": self.text, "chunks": self.chunks,
            "chunks_ready": self.chunks_ready, "chunk_urls": [f"{base}/chunks/{i}" for i in range(self.chunks_ready)],
            "duration": round(self.duration, 2), "audio_url": f"{base}/audio" if self.final else None,
            "format": self.format, "voice": self.voice, "lang": self.lang, "truncated": self.truncated,
            "queue_position": QUEUE_POS.get(self.id, 0), "cached": self.cached, "pause": self.pause,
        }


JOBS: Dict[str, Job] = {}
QUEUE: "asyncio.Queue[Job]" = asyncio.Queue(maxsize=QUEUE_MAX_SIZE)
QUEUE_POS: Dict[str, int] = {}


def _wav_bytes(wav: np.ndarray, sr: int) -> bytes:
    return encode_audio(wav, sr, "wav")


def _cache_key(job: Job) -> str:
    h = hashlib.sha256()
    h.update(json.dumps({"t": job.text, "v": job.voice, "l": job.lang, "s": job.speed, "p": job.pause,
                         "f": job.format, "gap": CHUNK_GAP_SECONDS, "ver": 3, "model": MODEL,
                         "bitrate": MP3_BITRATE, "revision": CACHE_REVISION, "engine": ENGINE_VERSION,
                         "sample_rate": getattr(job, "sample_rate", None),
                         "normalizer": hashlib.sha256((BASE_DIR / "textnorm.py").read_bytes()).hexdigest()}, ensure_ascii=False).encode())
    return h.hexdigest()[:32]


def _restore_from_cache(job: Job, key: str) -> bool:
    src = CACHE_DIR / key
    meta = src / "meta.json"
    if not meta.exists():
        return False
    try:
        m = json.loads(meta.read_text())
        if m["format"] not in OUTPUT_FORMATS or not 0 <= m["n"] <= MAX_TEXT_CHARS:
            return False
        job.check_cancel()
        for i in range(m["n"]):
            shutil.copyfile(src / f"chunk{i}.wav", job.dir / f"chunk{i}.wav")
        final_name = f"full.{m['format']}"
        shutil.copyfile(src / final_name, job.dir / final_name)
        os.utime(src, None)  # renova o TTL do cache
    except Cancelled:
        raise
    except Exception as e:
        logger.warning("cache corrompido %s: %s", key, e)
        shutil.rmtree(src, ignore_errors=True)
        return False
    job.check_cancel()
    job.chunks, job.chunks_ready, job.duration = m["chunks"], m["n"], m["duration"]
    job.format, job.final, job.cached = m["format"], str(job.dir / final_name), True
    job.status, job.stage, job.percent, job.message = "done", "done", 100, "Pronto (do cache)"
    return True


def _save_to_cache(job: Job, key: str) -> None:
    dst = CACHE_DIR / key
    tmp = Path(tempfile.mkdtemp(prefix=".publish-", dir=CACHE_DIR))
    try:
        for i in range(len(job.chunks)):
            job.check_cancel()
            shutil.copyfile(job.dir / f"chunk{i}.wav", tmp / f"chunk{i}.wav")
        shutil.copyfile(job.final, tmp / f"full.{job.format}")
        (tmp / "meta.json").write_text(json.dumps({"n": len(job.chunks), "chunks": job.chunks,
            "duration": job.duration, "format": job.format}, ensure_ascii=False))
        job.check_cancel()
        try:
            tmp.rename(dst)  # atomic publication; never delete a live cache entry
        except OSError:
            if not dst.is_dir():
                raise
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def _run_job_sync(job: Job, state) -> None:
    job.check_cancel()
    sr = state.tts.sample_rate
    job.sample_rate = sr
    inp = job.input
    # 1) texto
    if inp.get("url"):
        with tempfile.TemporaryDirectory() as tmp:
            media = _download_media(inp["url"], tmp, job.progress)
            text = _transcribe(media, job.lang, job.progress)
    elif inp.get("file"):
        name, data = inp["file"]
        text = _extract_from_upload(name, data, job.lang, job.progress, job.pages)
    else:
        text = clean_layout(inp.get("text") or "")
    job.check_cancel()
    if not text:
        raise UsarError("Não encontrei texto para ler.", 422, "empty_text")
    if len(text) > MAX_TEXT_CHARS:
        text, job.truncated = text[:MAX_TEXT_CHARS], True

    # Processamento opcional com IA (Tradução, Resumo, Podcast) e Limpeza Inteligente
    import ai_features
    from textnorm import smart_skip_filter
    if getattr(job, "smart_skip", False):
        text = smart_skip_filter(text)
        job.check_cancel()

    if getattr(job, "translate", False):
        job.progress("read", 25, "Traduzindo documento para português…")
        text = ai_features.translate_to_portuguese(text)
        job.check_cancel()

    mode = getattr(job, "mode", "normal")
    if mode == "summary":
        job.progress("read", 35, "Gerando resumo executivo do documento…")
        text = ai_features.summarize_text(text)
        job.check_cancel()
    elif mode.startswith("podcast"):
        style = mode.split("-", 1)[1] if "-" in mode else "fun"
        job.progress("read", 35, f"Criando roteiro de podcast ({style})…")
        script = ai_features.generate_podcast_script(text, style=style)
        text = script
        job.check_cancel()
        podcast_items = ai_features.parse_podcast_script(script, default_voice=job.voice)
        job.chunks = [t for _, t in podcast_items]
        job.podcast_voices = [v for v, _ in podcast_items]

    job.text = text
    if getattr(job, "transcribe_only", False):
        job.status, job.stage, job.percent, job.message = "done", "done", 100, "Pronto"
        return

    # 2) cache
    key = _cache_key(job)
    cacheable = job.voice not in state.custom_styles and getattr(job, "mode", "normal") == "normal"
    if cacheable and _restore_from_cache(job, key):
        return

    if not getattr(job, "chunks", None):
        job.chunks = split_chunks(text)
    job.progress("tts", 50, "Gerando voz…")

    # Stream each synthesized block to disk; only one audio block stays in RAM.
    total = len(job.chunks)
    wav_path = job.dir / "full.wav"
    is_podcast = getattr(job, "mode", "").startswith("podcast")
    with sf.SoundFile(str(wav_path), mode="w", samplerate=sr, channels=1, subtype="PCM_16") as full:
        if is_podcast:
            jingle_intro = ai_features.generate_podcast_jingle(sr, is_intro=True)
            full.write(jingle_intro)
            full.write(np.zeros(int(sr * 0.35), dtype=np.float32))
            job.duration += 2.55

        for i, chunk in enumerate(job.chunks):
            job.check_cancel()
            spoken = normalize_for_tts(chunk, job.lang).strip() or chunk.strip()
            voice_to_use = job.podcast_voices[i] if getattr(job, "podcast_voices", None) and i < len(job.podcast_voices) else job.voice
            try:
                wav, dur = _do_synthesize(state, text=spoken, voice=voice_to_use, lang=job.lang or None,
                    speed=job.speed, steps=None, max_chunk_length=None, silence_duration=None)
            except UnknownVoice as e:
                raise UsarError(f"Voz desconhecida: {e}", 400, "unknown_voice") from e
            job.check_cancel()
            wav = (wav.squeeze(0) if wav.ndim == 2 else wav).astype(np.float32)
            gap = (job.pause if chunk.endswith("\n\n") else CHUNK_GAP_SECONDS) if i < total - 1 else 0
            tmp = job.dir / f"chunk{i}.part"
            with sf.SoundFile(str(tmp), mode="w", samplerate=sr, channels=1, format="WAV", subtype="PCM_16") as block:
                block.write(wav)
                full.write(wav)
                silence = np.zeros(int(sr * gap), dtype=np.float32)
                block.write(silence)
                full.write(silence)
            tmp.replace(job.dir / f"chunk{i}.wav")
            job.duration += dur + gap
            job.chunks_ready = i + 1
            job.progress("tts", 50 + int(45 * (i + 1) / total), f"Gerando voz… bloco {i + 1} de {total}")

        if is_podcast:
            full.write(np.zeros(int(sr * 0.25), dtype=np.float32))
            jingle_outro = ai_features.generate_podcast_jingle(sr, is_intro=False)
            full.write(jingle_outro)
            job.duration += 2.05
    job.progress("encode", 96, "Montando o arquivo…")
    fmt = job.format
    out = job.dir / f"full.{fmt}"
    if fmt == "mp3":
        try:
            run_process([FFMPEG, "-y", "-loglevel", "error", "-i", str(wav_path), "-codec:a",
                         "libmp3lame", "-b:a", MP3_BITRATE, "-ac", "1", str(out)], job.check_cancel)
        except Cancelled:
            raise
        except Exception as e:
            logger.warning("mp3 falhou (%s); entregando wav", e)
            job.format, out = "wav", wav_path
    elif fmt == "m4b":
        try:
            from document_features import export_m4b, AudioChapter
            title = None
            inp = getattr(job, "input", {}) or {}
            if isinstance(inp, dict) and inp.get("file"):
                title = str(inp["file"][0]).rsplit(".", 1)[0]
            chapters = []
            if len(job.chunks) > 1:
                cur = 0.0
                step = max(0.5, job.duration / len(job.chunks))
                for idx, c in enumerate(job.chunks):
                    c_title = c.strip().split("\n")[0][:40] or f"Capítulo {idx + 1}"
                    chapters.append(AudioChapter(title=c_title, start=cur, end=min(job.duration, cur + step)))
                    cur += step
            export_m4b(wav_path, out, chapters=chapters if chapters else None, title=title or "SuperTonic Audiolivro", author="SuperTonic")
        except Cancelled:
            raise
        except Exception as e:
            logger.warning("m4b falhou (%s); entregando mp3", e)
            job.format, out = ("mp3", job.dir / "full.mp3") if "mp3" in OUTPUT_FORMATS else ("wav", wav_path)
            if job.format == "mp3":
                run_process([FFMPEG, "-y", "-loglevel", "error", "-i", str(wav_path), "-codec:a",
                             "libmp3lame", "-b:a", MP3_BITRATE, "-ac", "1", str(out)], job.check_cancel)
    elif fmt != "wav":
        subtype = "VORBIS" if fmt == "ogg" else None
        with sf.SoundFile(str(wav_path)) as source, sf.SoundFile(str(out), mode="w", samplerate=sr,
                channels=1, format=fmt.upper(), subtype=subtype) as target:
            for block in source.blocks(blocksize=65536):
                job.check_cancel()
                target.write(block)
    job.check_cancel()
    job.final = str(out)
    if out != wav_path:
        wav_path.unlink(missing_ok=True)
    try:
        if cacheable:
            _save_to_cache(job, key)
    except Cancelled:
        raise
    except Exception as e:  # pragma: no cover
        logger.warning("não consegui salvar no cache: %s", e)
    job.check_cancel()
    job.status, job.stage, job.percent, job.message = "done", "done", 100, "Pronto"


async def _worker(state):
    while True:
        job = await QUEUE.get()
        QUEUE_POS.pop(job.id, None)
        for k in QUEUE_POS:
            QUEUE_POS[k] = max(1, QUEUE_POS[k] - 1)
        if job.status != "queued":
            QUEUE.task_done()
            continue
        job.status, job.stage, job.message = "running", "read", "Começando…"
        try:
            task = asyncio.create_task(asyncio.to_thread(_run_job_sync, job, state))
            try:
                await asyncio.shield(task)
                if job.status == "done" and job.final:
                    asyncio.create_task(supabase_persistence.persist_completed_job(job))
            except asyncio.CancelledError:
                # to_thread cancellation does not stop its thread. Wait for cooperation
                # before unlinking uploads or letting the upstream model shut down.
                job.cancel_event.set()
                try:
                    await task
                except Exception:
                    pass
                job.status, job.message = "cancelled", "Cancelado"
                job.final, job.chunks_ready = None, 0
                raise
        except Cancelled:
            job.status, job.message = "cancelled", "Cancelado"
            job.final = None
            job.chunks_ready = 0
        except UsarError as e:
            job.status, job.error, job.message = "error", {"message": e.message, "code": e.code}, e.message
        except subprocess.TimeoutExpired:
            job.status, job.error, job.message = "error", {"message": "Demorou demais.", "code": "timeout"}, "Demorou demais."
        except Exception as e:  # pragma: no cover
            logger.exception("Job %s falhou", job.id)
            job.status, job.error, job.message = "error", {"message": f"Erro inesperado: {type(e).__name__}.", "code": "internal_error"}, "Erro inesperado."
        finally:
            if job.input.get("file"):
                Path(job.input["file"][1]).unlink(missing_ok=True)
            job.input = {}
            job.finished = time.time()
            job.completion.set()
            QUEUE.task_done()
            if job.status == "cancelled":
                shutil.rmtree(job.dir, ignore_errors=True)


async def _janitor():
    while True:
        await asyncio.sleep(600)
        now = time.time()
        for jid, job in list(JOBS.items()):
            if job.status in {"done", "error", "cancelled"} and job.finished is not None and now - job.finished > JOB_TTL_SECONDS:
                shutil.rmtree(job.dir, ignore_errors=True)
                JOBS.pop(jid, None)
        for d in CACHE_DIR.iterdir():
            if not d.name.startswith(".") and d.is_dir() and now - d.stat().st_mtime > CACHE_TTL_SECONDS:
                shutil.rmtree(d, ignore_errors=True)
        for d in SHARE_DIR.iterdir():
            meta = d / "meta.json"
            try:
                exp = json.loads(meta.read_text())["expires"] if meta.exists() else 0
            except Exception:
                exp = 0
            if now > exp:
                shutil.rmtree(d, ignore_errors=True)


# ---------------------------------------------------------------------------
# Rate limit por IP (janela deslizante de 1 h)
# ---------------------------------------------------------------------------
_HITS: Dict[str, List[float]] = {}


def _client_ip(request: Request) -> str:
    xff = request.headers.get("x-forwarded-for")
    if xff and request.client and trusted_peer(request.client.host, TRUSTED_PROXIES):
        # Peel trusted proxies from the right; a client-supplied left prefix
        # must not win when a trusted reverse proxy appends the true address.
        chain = [part.strip() for part in xff.split(",")]
        peer = request.client.host
        for address in reversed(chain):
            if not trusted_peer(peer, TRUSTED_PROXIES):
                break
            peer = address
        return peer
    return request.client.host if request.client else "?"


def _rate_limited(request: Request) -> Optional[int]:
    """Devolve segundos até liberar, ou None se pode seguir."""
    if RATE_LIMIT_PER_HOUR <= 0:
        return None
    ip = _client_ip(request)
    now = time.time()
    hits = [t for t in _HITS.get(ip, []) if now - t < 3600]
    if len(hits) >= RATE_LIMIT_PER_HOUR:
        _HITS[ip] = hits
        return int(3600 - (now - hits[0])) + 1
    hits.append(now)
    _HITS[ip] = hits
    if len(_HITS) > 5000:  # evita crescer sem limite
        for k in [k for k, v in _HITS.items() if not v or now - v[-1] > 3600][:1000]:
            _HITS.pop(k, None)
    return None


def _share_page(meta: Dict[str, Any], token: str) -> str:
    text = html.escape(meta.get("text", ""))
    title = html.escape(meta.get("title") or "Áudio do TrustVoice")
    dur = meta.get("duration", 0)
    mins, secs = int(dur // 60), int(dur % 60)
    exp = time.strftime("%d/%m/%Y %H:%M", time.localtime(meta["expires"]))
    return f"""<!DOCTYPE html><html lang="pt-BR"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{title} — TrustVoice</title><link rel="icon" href="/favicon.svg"><link rel="stylesheet" href="/static/styles.css">
<meta property="og:title" content="{title}"><meta property="og:description" content="Áudio de {mins}min {secs:02d}s gerado com TrustVoice">
</head><body><header class="top"><div class="brand"><img src="/favicon.svg" alt="" width="22" height="22"><span>TrustVoice</span></div><a class="link" href="/">Criar o meu</a></header>
<main><section class="hero"><h1>{title}</h1><p class="lead">{mins}min {secs:02d}s · {html.escape(meta.get("voice_label", ""))} · link válido até {exp}</p></section>
<section class="result"><audio controls preload="metadata" src="/s/{token}/audio"></audio>
<div class="result-actions"><a class="btn primary" href="/s/{token}/audio" download="{html.escape(meta.get("filename", "audio"))}">Baixar</a></div>
<div class="text" style="max-height:none">{text}</div></section></main>
<footer class="foot"><span>TrustVoice · Trust Corporation · Powered by Supertonic TTS</span><a href="/">trustvoice</a></footer></body></html>"""


# ---------------------------------------------------------------------------
# App
# ---------------------------------------------------------------------------
def build_app() -> FastAPI:
    app = create_app(model=MODEL)
    state = app.state.server_state

    workers: Dict[str, Any] = {}
    original_lifespan = app.router.lifespan_context

    @asynccontextmanager
    async def lifespan(app):
        async with original_lifespan(app):
            try:
                yield
            finally:
                for job in JOBS.values():
                    if job.status in {"queued", "running"}:
                        job.cancel_event.set()
                for task in workers.values():
                    task.cancel()
                await asyncio.gather(*workers.values(), return_exceptions=True)
                while not QUEUE.empty():
                    job = QUEUE.get_nowait()
                    if job.input.get("file"):
                        Path(job.input["file"][1]).unlink(missing_ok=True)
                    job.input = {}
                    job.status, job.finished = "cancelled", time.time()
                    job.completion.set()
                    QUEUE_POS.pop(job.id, None)
                    shutil.rmtree(job.dir, ignore_errors=True)
                    QUEUE.task_done()
    app.router.lifespan_context = lifespan

    def _ensure_workers():
        # create_app já define um lifespan, então on_event("startup") não dispara;
        # iniciamos as tarefas na primeira requisição que precisa delas.
        if not workers:
            workers["worker"] = asyncio.create_task(_worker(state))
            workers["janitor"] = asyncio.create_task(_janitor())

    @app.get("/", include_in_schema=False)
    async def root():
        return FileResponse(STATIC_DIR / "index.html")

    @app.get("/health", include_in_schema=False)
    async def health():
        ready = bool(getattr(state, "is_ready", False) and state.tts is not None)
        return {
            "status": "ok" if ready else "loading", "model": MODEL, "service": "supertonic",
            "document": _has("pypdf") and _has("docx"),
            "ocr": _has("pytesseract") and shutil.which("tesseract") is not None,
            "video": bool(FFMPEG) and (bool(OPENAI_KEY) or _has("faster_whisper")),
            "transcriber": "openai" if OPENAI_KEY else ("local" if _has("faster_whisper") else None),
            "remote_downloads": ALLOW_REMOTE_DOWNLOADS, "mp3": "mp3" in OUTPUT_FORMATS, "queue": QUEUE.qsize(), "rate_limit_per_hour": RATE_LIMIT_PER_HOUR,
            "share_ttl_hours": SHARE_TTL_SECONDS // 3600, "cache": sum(1 for d in CACHE_DIR.iterdir() if d.is_dir()),
        }

    @app.get("/api/voices", include_in_schema=False)
    async def voices():
        names = list(state.tts.voice_style_names) if state.tts else []
        return {"voices": names, "custom": sorted(state.custom_styles.keys()), "default_lang": DEFAULT_LANG,
                "formats": OUTPUT_FORMATS, "max_upload_mb": MAX_UPLOAD_MB, "max_text_chars": MAX_TEXT_CHARS}

    async def _read_inputs(text, url, file):
        if url and url.strip():
            if not ALLOW_REMOTE_DOWNLOADS:
                raise UsarError("Downloads remotos desabilitados. Envie o arquivo.", 403, "remote_download_disabled")
            return {"url": url.strip()}
        if file is not None and file.filename:
            try:
                path = await save_upload(file, JOBS_DIR, MAX_UPLOAD_MB * 1024 * 1024)
            except ResourceError as e:
                code = str(e)
                raise UsarError("Arquivo vazio." if code == "empty_file" else "Arquivo muito grande.",
                                400 if code == "empty_file" else 413, code)
            return {"file": (file.filename, path)}
        if text and text.strip():
            return {"text": text}
        raise UsarError("Envie um texto, um arquivo ou um link.", 400, "missing_input")

    @app.post("/api/jobs", include_in_schema=False)
    async def create_job(
        request: Request,
        text: Optional[str] = Form(None), url: Optional[str] = Form(None), voice: str = Form("F1"),
        lang: Optional[str] = Form(None), speed: Optional[float] = Form(None),
        response_format: str = Form("mp3"), file: Optional[UploadFile] = File(None),
        pause: Optional[float] = Form(None), pages: Optional[str] = Form(None),
        mode: str = Form("normal"), translate: bool = Form(False), smart_skip: bool = Form(False),
    ):
        if state.tts is None:
            return _err(503, "Modelo ainda carregando. Tente em alguns segundos.", "model_loading", "server_error")
        wait = _rate_limited(request)
        if wait:
            return JSONResponse(status_code=429, headers={"Retry-After": str(wait)}, content={"error": {
                "message": f"Muitas gerações neste IP. Tente de novo em {max(1, wait // 60)} min.",
                "type": "rate_limit_error", "code": "rate_limited"}})
        fmt = (response_format or "mp3").lower()
        if fmt == "mp3" and "mp3" not in OUTPUT_FORMATS:
            fmt = "wav"
        if fmt not in OUTPUT_FORMATS:
            return _err(400, f"Formato inválido. Use: {', '.join(OUTPUT_FORMATS)}.", "unsupported_response_format")
        try:
            inputs = await _read_inputs(text, url, file)
        except UsarError as e:
            return _err(e.status, e.message, e.code)
        _ensure_workers()
        pause_s = DEFAULT_PARAGRAPH_PAUSE if pause is None else max(0.0, min(3.0, float(pause)))
        mode_str = str(getattr(mode, "default", mode) if hasattr(mode, "default") else (mode or "normal"))
        translate_bool = bool(getattr(translate, "default", translate) if hasattr(translate, "default") else translate)
        smart_skip_bool = bool(getattr(smart_skip, "default", smart_skip) if hasattr(smart_skip, "default") else smart_skip)
        job = Job(voice=voice or "F1", lang=(lang or DEFAULT_LANG), speed=speed, format=fmt, pause=pause_s,
                  pages=(pages or "").strip() or None, mode=mode_str, translate=translate_bool, smart_skip=smart_skip_bool, **inputs)
        job.transcribe_only = bool(inputs.get("url") and request.url.path == "/usar")
        try:
            QUEUE.put_nowait(job)
        except asyncio.QueueFull:
            if inputs.get("file"):
                Path(inputs["file"][1]).unlink(missing_ok=True)
            shutil.rmtree(job.dir, ignore_errors=True)
            return _err(429, "Fila cheia. Tente novamente depois.", "queue_full")
        JOBS[job.id] = job
        QUEUE_POS[job.id] = QUEUE.qsize()
        return JSONResponse(job.public(request), status_code=202)

    @app.get("/api/jobs/{job_id}", include_in_schema=False)
    async def get_job(job_id: str, request: Request):
        job = JOBS.get(job_id)
        if not job:
            if supabase_persistence.is_enabled():
                meta = supabase_persistence.get_job_meta(job_id)
                if meta:
                    return meta
            return _err(404, "Job não encontrado (expirou?).", "not_found")
        return job.public(request)

    @app.get("/api/jobs/{job_id}/chunks/{n}", include_in_schema=False)
    async def get_chunk(job_id: str, n: int):
        job = JOBS.get(job_id)
        if not job or n < 0 or n >= job.chunks_ready:
            return _err(404, "Bloco não disponível.", "not_found")
        return FileResponse(job.dir / f"chunk{n}.wav", media_type="audio/wav", headers={"Cache-Control": "private, max-age=86400"})

    @app.get("/api/jobs/{job_id}/audio", include_in_schema=False)
    async def get_audio(job_id: str):
        job = JOBS.get(job_id)
        if job and job.status == "done" and job.final and Path(job.final).exists():
            mime = "audio/mpeg" if job.format == "mp3" else format_to_mime(job.format)
            return FileResponse(job.final, media_type=mime, filename=f"trustvoice-{job.id}.{job.format}",
                                headers={"Cache-Control": "private, max-age=86400"})
        if supabase_persistence.is_enabled():
            fmt = job.format if job else "mp3"
            signed_url = supabase_persistence.get_audio_signed_url(job_id, fmt)
            if signed_url:
                from starlette.responses import RedirectResponse
                return RedirectResponse(signed_url, status_code=307)
        return _err(404, "Áudio ainda não está pronto.", "not_ready")

    @app.post("/api/jobs/{job_id}/share", include_in_schema=False)
    async def share_job(job_id: str, request: Request, title: Optional[str] = Form(None)):
        job = JOBS.get(job_id)
        if not job or job.status != "done" or not job.final:
            return _err(404, "Áudio ainda não está pronto.", "not_ready")
        token = secrets.token_urlsafe(9)
        d = SHARE_DIR / token
        d.mkdir(parents=True, exist_ok=True)
        ext = job.format
        shutil.copyfile(job.final, d / f"audio.{ext}")
        expires = time.time() + SHARE_TTL_SECONDS
        meta = {"text": job.text, "title": (title or "").strip()[:120] or None, "duration": job.duration, "format": ext,
                "voice_label": job.voice, "filename": f"trustvoice-{token}.{ext}", "expires": expires}
        (d / "meta.json").write_text(json.dumps(meta, ensure_ascii=False))
        base = str(request.base_url).rstrip("/")
        if request.client and trusted_peer(request.client.host, TRUSTED_PROXIES) and request.headers.get("x-forwarded-proto") == "https":
            base = "https://" + base.split("://", 1)[1]
        return {"url": f"{base}/s/{token}", "expires": expires}

    def _share_meta(token: str) -> Optional[Dict[str, Any]]:
        if not re.fullmatch(r"[A-Za-z0-9_-]{6,32}", token):
            return None
        meta = SHARE_DIR / token / "meta.json"
        if not meta.exists():
            return None
        m = json.loads(meta.read_text())
        if time.time() > m["expires"]:
            shutil.rmtree(SHARE_DIR / token, ignore_errors=True)
            return None
        return m

    @app.get("/s/{token}", include_in_schema=False)
    async def share_page(token: str):
        m = _share_meta(token)
        if not m:
            return Response("<h1>Link expirado ou inválido.</h1><p><a href='/'>Voltar ao TrustVoice</a></p>",
                            status_code=404, media_type="text/html")
        return Response(_share_page(m, token), media_type="text/html")

    @app.get("/s/{token}/audio", include_in_schema=False)
    async def share_audio(token: str):
        m = _share_meta(token)
        if not m:
            return _err(404, "Link expirado ou inválido.", "not_found")
        mime = "audio/mpeg" if m["format"] == "mp3" else format_to_mime(m["format"])
        return FileResponse(SHARE_DIR / token / f"audio.{m['format']}", media_type=mime,
                            headers={"Cache-Control": "public, max-age=3600"})

    @app.delete("/api/jobs/{job_id}", include_in_schema=False)
    async def delete_job(job_id: str):
        job = JOBS.get(job_id)
        if job:
            if job.status in {"queued", "running"}:
                job.cancel_event.set()
                job.message = "Cancelamento solicitado…"
        return {"ok": True}

    # ---- compatibilidade: rota síncrona antiga ------------------------------
    @app.post("/usar", include_in_schema=False)
    async def usar(
        text: Optional[str] = Form(None), url: Optional[str] = Form(None), voice: str = Form("F1"),
        lang: Optional[str] = Form(None), speed: Optional[float] = Form(None),
        response_format: str = Form("wav"), file: Optional[UploadFile] = File(None),
        pages: Optional[str] = Form(None), request: Request = None,
    ):
        if state.tts is None:
            return _err(503, "Modelo ainda carregando.", "model_loading", "server_error")
        fmt = (response_format or "wav").lower()
        if fmt not in SUPPORTED_FORMATS:
            return _err(400, f"Formato inválido. Use: {', '.join(SUPPORTED_FORMATS)}.", "unsupported_response_format")
        lang = lang or DEFAULT_LANG
        # Use the same bounded worker and streaming implementation as /api/jobs.
        result = await create_job(request=request, text=text, url=url, voice=voice, lang=lang,
            speed=speed, response_format=fmt, file=file, pause=None, pages=pages)
        if result.status_code != 202:
            return result
        job = JOBS[json.loads(result.body)["id"]]
        try:
            await job.completion.wait()
        except asyncio.CancelledError:
            job.cancel_event.set()
            raise
        if job.status != "done":
            error = job.error or {"message": "Cancelado", "code": "cancelled"}
            return _err(422, error["message"], error["code"])
        if getattr(job, "transcribe_only", False):
            return JSONResponse({"text": job.text, "truncated": job.truncated})
        return FileResponse(job.final, media_type=format_to_mime(job.format), headers={
            "X-Texto": urllib.parse.quote(job.text[:8000]), "X-Duration-Seconds": f"{job.duration:.2f}",
            "X-Truncated": "1" if job.truncated else "0", "Cache-Control": "no-store"})

    @app.get("/api/documents", include_in_schema=False)
    async def get_saved_documents():
        if not supabase_persistence.is_enabled():
            return JSONResponse({"enabled": False, "documents": []})
        docs = await asyncio.to_thread(supabase_persistence.list_saved_documents)
        return JSONResponse({"enabled": True, "documents": docs})

    @app.get("/api/documents/{job_uuid}/audio", include_in_schema=False)
    async def get_saved_document_audio(job_uuid: str):
        if not supabase_persistence.is_enabled():
            return _err(404, "Supabase não configurado.", "not_configured")
        signed_url = await asyncio.to_thread(supabase_persistence.get_audio_url_by_uuid, job_uuid)
        if not signed_url:
            return _err(404, "Áudio não encontrado no Supabase.", "not_found")
        from starlette.responses import RedirectResponse
        return RedirectResponse(signed_url, status_code=307)

    @app.delete("/api/documents/{job_uuid}", include_in_schema=False)
    async def delete_saved_document(job_uuid: str):
        if not supabase_persistence.is_enabled():
            return _err(404, "Supabase não configurado.", "not_configured")
        ok = await asyncio.to_thread(supabase_persistence.delete_saved_document, job_uuid)
        return JSONResponse({"success": ok})

    @app.get("/api/feed.xml", include_in_schema=False)
    @app.get("/feed.xml", include_in_schema=False)
    async def podcast_rss_feed(request: Request):
        import ai_features
        base_url = str(request.base_url).rstrip("/")
        if request.client and trusted_peer(request.client.host, TRUSTED_PROXIES) and request.headers.get("x-forwarded-proto") == "https":
            base_url = "https://" + base_url.split("://", 1)[1]
        docs = await asyncio.to_thread(supabase_persistence.list_saved_documents)
        xml_content = ai_features.build_podcast_rss_feed(docs, base_url)
        return Response(content=xml_content, media_type="application/rss+xml; charset=utf-8")

    @app.get("/manifest.webmanifest", include_in_schema=False)
    async def manifest():
        return FileResponse(STATIC_DIR / "manifest.webmanifest", media_type="application/manifest+json")

    @app.get("/sw.js", include_in_schema=False)
    async def sw():
        return FileResponse(STATIC_DIR / "sw.js", media_type="application/javascript")

    @app.get("/favicon.svg", include_in_schema=False)
    async def favicon():
        return FileResponse(STATIC_DIR / "favicon.svg", media_type="image/svg+xml")

    app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")

    api_key = (os.environ.get("API_KEY") or "").strip()
    if api_key:
        app.add_middleware(ApiKeyMiddleware, api_key=api_key)
    else:
        logger.warning("API_KEY não definida — /v1/* está aberto.")
    app.add_middleware(BodyLimitMiddleware, max_bytes=MAX_UPLOAD_MB * 1024 * 1024 + 1024 * 1024)
    return app


app = build_app()


def main() -> None:
    uvicorn.run("server:app", host=os.environ.get("HOST", "0.0.0.0"), port=int(os.environ.get("PORT", "8080")),
                proxy_headers=False, log_level=os.environ.get("LOG_LEVEL", "info"))


if __name__ == "__main__":
    main()
